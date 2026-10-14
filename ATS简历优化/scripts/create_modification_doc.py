#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
生成修改说明文档
"""

import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

os.makedirs('./简历优化', exist_ok=True)

def create_modification_doc(direction, color_theme):
    """创建修改说明文档"""
    doc = Document()
    
    # 设置默认字体
    style = doc.styles['Normal']
    style.font.name = '微软雅黑'
    style.font.size = Pt(10.5)
    
    # 标题
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f'贾小小简历优化 - {direction}方向')
    run.bold = True
    run.font.size = Pt(18)
    
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('详细修改说明文档')
    run.bold = True
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(*color_theme)
    
    doc.add_paragraph()
    
    # ===== 一、整体优化思路 =====
    p = doc.add_paragraph()
    run = p.add_run('一、整体优化思路与核心逻辑')
    run.bold = True
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor(*color_theme)
    
    if "数字化HR" in direction:
        strategies = [
            '聚焦HR数字化转型能力：将简历中涉及系统、工具、数据、流程的内容前置并强化',
            '突出数字化关键词：HRIS、HRSSC、数字化转型、数据驱动、流程自动化等',
            '强调技术应用能力：AI+HR、数据分析、系统实施等技能凸显',
            '弱化传统HR表述：减少纯管理类描述，强化技术+业务复合能力',
            '成果数据化呈现：将所有工作成果用数字化指标呈现',
        ]
    else:
        strategies = [
            '聚焦全盘管理能力：突出战略规划、组织发展、全模块统筹能力',
            '强调管理幅度与层级：凸显400-500人团队、十多家分子公司统筹经验',
            '突出业务伙伴角色：展现与高管层的战略对话能力',
            '强化团队建设表述：展示HR团队组建、领导、培养经验',
            '弱化技术细节：相对弱化数字化工具，聚焦管理价值输出',
        ]
    
    for s in strategies:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.left_indent = Inches(0.3)
        p.add_run(s)
    
    doc.add_paragraph()
    
    # ===== 二、分模块具体修改点 =====
    p = doc.add_paragraph()
    run = p.add_run('二、分模块具体修改点及修改原因')
    run.bold = True
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor(*color_theme)
    
    modifications = [
        ('求职意向', '明确写出目标岗位名称，数字化HR方向写"数字化人力资源管理 | HR数字化转型 | HRIS"，全盘HR方向写"人力资源总监 | 全盘HR管理负责人 | HRM/HRD"'),
        ('工作经历排序', '保持时间倒序，将最相关经历放在前面，重点突出奇迹科技和现公司经历'),
        ('职责描述优化', '统一采用「执行动作 + 工作内容 + 可量化成果」句式，每段经历保留6-8条核心职责'),
        ('关键词植入', '根据目标岗位JD要求，在各段经历中自然融入行业关键词'),
        ('数据呈现强化', '所有量化成果保留并前置，如"招聘周期缩短30%"、"事务自动化90%"等'),
    ]
    
    for title, desc in modifications:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.2)
        run = p.add_run(f'{title}：')
        run.bold = True
        p.add_run(desc)
    
    doc.add_paragraph()
    
    # ===== 三、方向差异化调整 =====
    p = doc.add_paragraph()
    run = p.add_run('三、两个方向的差异化调整')
    run.bold = True
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor(*color_theme)
    
    table = doc.add_table(rows=5, cols=3)
    table.style = 'Table Grid'
    
    # 表头
    headers = ['模块', '数字化HR方向', '全盘HR方向']
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = header
        cell.paragraphs[0].runs[0].bold = True
    
    # 内容
    rows_data = [
        ('求职意向', '数字化人力资源管理 | HRIS', '人力资源总监 | 全盘HR管理'),
        ('核心亮点', 'HR数字化转型、HRSSC建设', '集团统筹、战略伙伴角色'),
        ('关键词侧重', 'HRIS/HR SaaS、流程自动化', '战略规划、组织设计、人才梯队'),
        ('成果呈现', '系统上线、效率提升百分比', '团队规模、管理幅度、业务支撑'),
    ]
    
    for i, (col1, col2, col3) in enumerate(rows_data, 1):
        table.rows[i].cells[0].text = col1
        table.rows[i].cells[1].text = col2
        table.rows[i].cells[2].text = col3
    
    doc.add_paragraph()
    
    # ===== 四、关键词优化依据 =====
    p = doc.add_paragraph()
    run = p.add_run('四、关键词优化依据和排版调整考量')
    run.bold = True
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor(*color_theme)
    
    if "数字化HR" in direction:
        keywords = [
            'HR数字化转型：行业热点，体现技术敏感度',
            'HRIS/HRSSC：核心技术能力，直接命中岗位要求',
            '流程自动化/RPA：体现技术应用能力',
            '人力资本分析（HCA）：体现数据驱动思维',
            '智能招聘/ATS：体现数字化工具应用',
            '招聘周期缩短30%：量化成果，增强说服力',
        ]
    else:
        keywords = [
            '人力资源战略规划：体现战略高度',
            '组织架构设计与优化：核心管理能力',
            '人才梯队建设：长期人才保障能力',
            '薪酬绩效体系设计：专业能力体现',
            'HR战略伙伴角色：业务导向思维',
            '集团化管理：大型组织统筹能力',
        ]
    
    for kw in keywords:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.left_indent = Inches(0.3)
        p.add_run(kw)
    
    doc.add_paragraph()
    
    # ===== 五、被调整内容标注 =====
    p = doc.add_paragraph()
    run = p.add_run('五、被调整/合并/补充内容标注')
    run.bold = True
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor(*color_theme)
    
    adjustments = [
        ('个人优势模块', '原简历有独立"个人优势"模块，优化版将其内容整合到工作经历的职责描述中，通过内容排序和表述方式隐性体现'),
        ('工作经历细化', '部分工作经历进行了职责合并或拆分，确保每段经历6-8条核心职责，避免信息冗余'),
        ('数据保留', '所有原始简历中的量化数据（如招聘周期缩短30%、员工满意度80%等）均已保留，仅调整表述方式'),
        ('时间线调整', '保持时间倒序，每段经历起始时间与原始简历一致'),
        ('企业简介精简', '部分公司的企业简介进行了精简处理，保留核心信息，节省篇幅'),
    ]
    
    for title, desc in adjustments:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.2)
        run = p.add_run(f'{title}：')
        run.bold = True
        p.add_run(desc)
    
    doc.add_paragraph()
    
    # ===== 六、后续微调方法 =====
    p = doc.add_paragraph()
    run = p.add_run('六、后续个性化微调方法')
    run.bold = True
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor(*color_theme)
    
    tips = [
        '针对具体企业：在现有简历基础上，根据目标企业的JD关键词进行针对性调整',
        '突出特定经历：根据不同企业背景，将最相关的经历放在更显眼的位置',
        '调整关键词密度：确保目标岗位的核心关键词在简历中出现2-3次',
        '量化成果强化：根据目标企业的规模和问题，适当调整成果数据的呈现方式',
        '排版微调：根据不同平台要求（如Boss直聘、猎聘等），调整简历篇幅和排版',
    ]
    
    for tip in tips:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.left_indent = Inches(0.3)
        p.add_run(tip)
    
    # 保存
    filename = f'./简历优化/贾小小_{direction}_修改说明.docx'
    doc.save(filename)
    print(f'{direction}修改说明文档已保存')

# 执行
create_modification_doc('数字化HR', (102, 126, 234))
create_modification_doc('全盘HR', (245, 87, 108))
print('修改说明文档生成完成！')
